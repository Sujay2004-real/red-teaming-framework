import asyncio
import io
from datetime import timedelta
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, text, inspect
from sqlalchemy.orm import sessionmaker

import main
from database import Base, Assessment, Target, ToolExecution, ExecutionAttempt, Finding, FindingReview, get_db, utcnow
from modules.jobs import run_execution, pending_jobs
from modules.scope_rules import scope_contains, scopes_overlap, command_destinations
from modules.outcomes import outcome
from conftest import OPERATOR_TEST_KEY


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    engine = create_engine('sqlite:///' + str(tmp_path / 'test.db'), connect_args={'check_same_thread': False})
    @event.listens_for(engine, 'connect')
    def setup(connection, _):
        connection.execute('PRAGMA foreign_keys=ON')
        connection.execute('PRAGMA busy_timeout=10000')
    Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, expire_on_commit=False)
    def database():
        with sessions() as db:
            yield db
    main.app.dependency_overrides[get_db] = database
    monkeypatch.setattr(main, 'REPORTS_DIR', tmp_path / 'reports')
    monkeypatch.setattr(main, 'schedule_auto_recommendations', lambda *args, **kwargs: None)
    client = TestClient(main.app, headers={'X-API-Key': OPERATOR_TEST_KEY})
    yield client, sessions
    main.app.dependency_overrides.clear()
    engine.dispose()


def create(client, policy=None):
    payload = {'name': 'Lab', 'scope_domain_ip': 'lab.test', 'authorized_scopes': ['lab.test']}
    if policy is not None:
        payload['engagement_policy'] = policy
    target = client.post('/targets/', json=payload).json()
    response = client.post('/assessments/', json={'target_id': target['id'], 'objective': 'Inspect',
        'plan': [{'tool': 'curl', 'command': 'curl -I http://lab.test'}]})
    assert response.status_code == 200
    return response.json()


def queue(client, assessment):
    response = client.post(f"/assessments/{assessment['id']}/execute", json={
        'step_index': 0, 'approved': True, 'background': True, 'plan_version': assessment['plan_version']})
    assert response.status_code == 202, response.text
    return response.json()['execution_id']


def test_queue_runs_once_and_preserves_attempts(workspace, monkeypatch):
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    fake = AsyncMock(return_value={'stdout': 'first attempt', 'stderr': 'failed', 'return_code': 7, 'duration_ms': 4})
    monkeypatch.setattr(main.executor, 'execute_command', fake)
    async def concurrent():
        return await asyncio.gather(run_execution(execution_id, sessions), run_execution(execution_id, sessions))
    asyncio.run(concurrent())
    assert fake.await_count == 1
    queue(client, assessment)
    fake.return_value = {'stdout': 'second attempt', 'stderr': '', 'return_code': 0, 'duration_ms': 5}
    asyncio.run(run_execution(execution_id, sessions))
    history = client.get(f"/assessments/{assessment['id']}/executions/{execution_id}/attempts").json()
    assert [a['stdout'] for a in history] == ['first attempt', 'second attempt']
    assert [a['attempt'] for a in history] == [1, 2]
    assert client.delete(f"/assessments/{assessment['id']}").status_code == 200


def test_cancel_queued_never_calls_executor(workspace, monkeypatch):
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    prefix = f"/assessments/{assessment['id']}/executions/{execution_id}"
    assert client.post(prefix + '/cancel').json()['state'] == 'cancelled'
    fake = AsyncMock()
    monkeypatch.setattr(main.executor, 'execute_command', fake)
    assert asyncio.run(run_execution(execution_id, sessions)) is None
    fake.assert_not_called()
    assert client.get(prefix + '/attempts').json()[0]['state'] == 'cancelled'


def test_cancel_running_persists_partial_output(workspace, monkeypatch):
    from modules.executor import live_registry
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    async def scanning(*args, **kwargs):
        live_registry.append(execution_id, 'partial evidence')
        with sessions() as db:
            db.get(ToolExecution, execution_id).cancel_requested = True
            db.commit()
        await asyncio.sleep(20)
    monkeypatch.setattr(main.executor, 'execute_command', scanning)
    result = asyncio.run(run_execution(execution_id, sessions))
    assert result['outcome'] == 'cancelled'
    with sessions() as db:
        row = db.get(ToolExecution, execution_id)
        assert row.stdout == 'partial evidence'
        assert db.query(ExecutionAttempt).count() == 1


def test_dead_worker_is_interrupted_without_reexecution(workspace):
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    with sessions() as db:
        row = db.get(ToolExecution, execution_id)
        row.state, row.worker_id = 'running', 'dead-worker'
        row.heartbeat_at = utcnow() - timedelta(seconds=61)
        db.commit()
    assert pending_jobs(sessions, 2) == []
    with sessions() as db:
        assert db.get(ToolExecution, execution_id).state == 'interrupted'
        assert db.get(Assessment, assessment['id']).status != 'running'
        assert db.query(ExecutionAttempt).count() == 1


def test_live_cursor_reads_persisted_output(workspace):
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    with sessions() as db:
        row = db.get(ToolExecution, execution_id)
        row.live_output, row.output_cursor = 'abcdef', 12
        db.commit()
    prefix = f"/assessments/{assessment['id']}/executions/{execution_id}/live"
    assert client.get(prefix + '?cursor=8').json()['output'] == 'cdef'
    assert client.get(prefix + '?cursor=0').json()['reset'] is True
    assert client.get(prefix + '?cursor=12').json()['output'] == ''


def test_stale_plan_and_changed_executed_phase_are_refused(workspace):
    client, _ = workspace
    assessment = create(client)
    prefix = f"/assessments/{assessment['id']}"
    changed = client.put(prefix + '/plan', json={'plan': assessment['plan'], 'plan_version': 1})
    assert changed.json()['plan_version'] == 2
    assert client.post(prefix + '/execute', json={'step_index': 0, 'approved': True, 'plan_version': 1}).status_code == 409
    queue(client, changed.json())
    altered = [{**assessment['plan'][0], 'phase': 'exploitation'}]
    assert client.put(prefix + '/plan', json={'plan': altered, 'plan_version': 2}).status_code == 409


@pytest.mark.parametrize('target,scope,allowed', [
    ('https://a.test/api/users', 'https://a.test/api', True),
    ('http://a.test/api', 'https://a.test/api', False),
    ('https://a.test:444/api', 'https://a.test/api', False),
    ('https://a.test/apievil', 'https://a.test/api', False),
    ('https://a.test/api/%2e%2e/admin', 'https://a.test/api', False),
    ('child.a.test', 'a.test', False), ('child.a.test', '*.a.test', True),
    ('10.0.0.0/24', '10.0.0.0/25', False),
])
def test_precise_scope_boundaries(target, scope, allowed):
    assert scope_contains(target, scope) is allowed


def test_exclusion_detects_ranges_and_resource_script_destinations():
    assert scopes_overlap('10.0.0.0/24', '10.0.0.128/25')
    assert command_destinations('msfconsole -x "use auxiliary/scanner/ftp/ftp_version; set RHOSTS 10.0.0.2; run"', []) == ['10.0.0.2']


@pytest.mark.parametrize('command', [
    'msfconsole -x "use auxiliary/scanner/ftp/ftp_version; set RHOSTS 10.0.0.2; run"',
    'msfconsole -x="use auxiliary/scanner/ftp/ftp_version; set RHOSTS 10.0.0.2; run"',
    'msfconsole --execute-command "use auxiliary/scanner/ftp/ftp_version; set RHOSTS 10.0.0.2; run"',
    'msfconsole --execute-command="use auxiliary/scanner/ftp/ftp_version; set RHOSTS 10.0.0.2; run"',
])
def test_msf_destinations_extracted_from_every_accepted_spelling(command):
    # The policy engine accepts the resource script under all four spellings
    # (policy_engine.validate_command reads it from -x and --execute-command,
    # attached or separated). command_destinations feeds the sole excluded_scopes
    # enforcement point in main.py, so it must surface RHOSTS from every one of
    # them; otherwise a spelling this parser missed would slip an in-range but
    # deliberately-excluded host past the exclusion check.
    assert command_destinations(command, []) == ['10.0.0.2']


def test_expired_policy_refused_at_approval(workspace):
    client, _ = workspace
    assessment = create(client, {'ends_at': '2020-01-01T00:00:00Z'})
    response = client.post(f"/assessments/{assessment['id']}/execute", json={'step_index': 0, 'approved': True})
    assert response.status_code == 403 and 'expired' in response.text


def test_summary_omits_evidence_and_review_requires_proof(workspace):
    client, sessions = workspace
    assessment = create(client)
    with sessions() as db:
        finding = Finding(assessment_id=assessment['id'], fingerprint='stable', title='Example', evidence='sensitive evidence', source_tools=['curl'], severity='Low')
        db.add(finding); db.commit(); fid = finding.id
    prefix = f"/assessments/{assessment['id']}"
    page = client.get(prefix + '/findings?limit=1').json()
    assert page['total'] == 1 and 'evidence' not in page['items'][0]
    assert 'plan' not in client.get('/catalog/assessments').json()['items'][0]
    assert client.put(prefix + f'/findings/{fid}/review', json={'status': 'resolved'}).status_code == 422
    assert client.put(prefix + f'/findings/{fid}/review', json={'status': 'accepted_risk'}).status_code == 422
    response = client.put(prefix + f'/findings/{fid}/review', json={'status': 'resolved', 'retest_evidence': 'Retest passed', 'comment': 'Reviewed'})
    assert response.status_code == 200 and len(response.json()['comments']) == 1
    assert client.post('/workspace/reset').status_code == 200


def test_report_versions_remain_after_analysis_invalidation(workspace, monkeypatch):
    client, sessions = workspace
    assessment = create(client)
    execution_id = queue(client, assessment)
    monkeypatch.setattr(main.executor, 'execute_command', AsyncMock(return_value={'stdout': 'HTTP/1.1 200 OK', 'stderr': '', 'return_code': 0, 'duration_ms': 5}))
    asyncio.run(run_execution(execution_id, sessions))
    prefix = f"/assessments/{assessment['id']}"
    assert client.post(prefix + '/analyze').status_code == 200
    report = client.post(prefix + '/report').json()
    before = client.get(report['download_url']).content
    assert b'Evidence provenance' in before
    with sessions() as db:
        main.invalidate_analysis(db, assessment['id']); db.commit()
    assert client.get(report['download_url']).content == before
    assert client.get(prefix + '/reports').json()[0]['digest'] == report['digest']


def test_docx_tables_and_extraction_limits():
    from docx import Document
    from modules.document_text import extract
    document = Document(); document.add_paragraph('Authorized targets')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'lab.test'; table.cell(0, 1).text = 'No exploitation'
    stream = io.BytesIO(); document.save(stream)
    assert 'lab.test | No exploitation' in extract('.docx', stream.getvalue())
    assert len(extract('.txt', b'x' * 40000)) == 30001


def test_zap_warning_is_successful_outcome():
    assert outcome('zap-baseline.py', 2) == 'completed_with_findings'
    assert outcome('nmap', 2) == 'failed'


def test_migrations_fresh_and_legacy(tmp_path):
    for legacy in (False, True):
        engine = create_engine('sqlite:///' + str(tmp_path / f'migration-{legacy}.db'))
        if legacy:
            with engine.begin() as connection:
                connection.execute(text('CREATE TABLE targets (id INTEGER PRIMARY KEY, name VARCHAR NOT NULL, scope_domain_ip VARCHAR NOT NULL, created_at DATETIME)'))
                connection.execute(text("INSERT INTO targets (id,name,scope_domain_ip) VALUES (1,'preserved','lab.test')"))
        config = Config(); config.set_main_option('script_location', str(Path(__file__).parent / 'migrations'))
        with engine.begin() as connection:
            config.attributes['connection'] = connection
            command.upgrade(config, 'head')
            command.upgrade(config, 'head')
            from database import migration_head
            assert connection.execute(text('SELECT version_num FROM alembic_version')).scalar() == migration_head()
            assert 'plan_version' in {c['name'] for c in inspect(connection).get_columns('assessments')}
            if legacy:
                assert connection.execute(text('SELECT name FROM targets WHERE id=1')).scalar() == 'preserved'
        engine.dispose()
