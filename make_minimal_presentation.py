"""Generate a concise project presentation."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "pptx" / "AI_Assisted_Red_Teaming_Framework_Minimal.pptx"

NAVY = RGBColor(18, 35, 58)
BLUE = RGBColor(34, 103, 172)
CYAN = RGBColor(39, 168, 181)
GREEN = RGBColor(46, 125, 94)
AMBER = RGBColor(191, 119, 33)
RED = RGBColor(173, 55, 47)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(32, 42, 54)
MUTED = RGBColor(94, 107, 120)
PALE = RGBColor(237, 243, 249)
LINE = RGBColor(197, 210, 223)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height


def add_slide(title=None, subtitle=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else WHITE
    if title:
        add_text(slide, .58, .27, 12.15, .48, title, 25, WHITE if dark else NAVY, True)
        add_shape(slide, .58, .9, 1.0, .07, CYAN)
    if subtitle:
        add_text(slide, 1.75, .68, 10.95, .3, subtitle, 11.5, RGBColor(191, 214, 232) if dark else MUTED)
    return slide


def add_shape(slide, x, y, w, h, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, value, size=15, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font="Aptos"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.03
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, items, size=15, color=INK, accent=CYAN):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = 1.08
        dot = p.add_run()
        dot.text = "●  "
        dot.font.name = "Aptos"
        dot.font.size = Pt(size - 2)
        dot.font.color.rgb = accent
        body = p.add_run()
        body.text = item
        body.font.name = "Aptos"
        body.font.size = Pt(size)
        body.font.color.rgb = color
    return tb


def card(slide, x, y, w, h, heading, body, color=BLUE, number=None, body_size=12.5):
    add_shape(slide, x, y, w, h, WHITE, LINE, True)
    add_shape(slide, x, y, .08, h, color, radius=True)
    if number is not None:
        add_shape(slide, x + .22, y + .22, .48, .48, color, radius=True)
        add_text(slide, x + .22, y + .22, .48, .48, str(number), 13, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        hx = x + .83
        hw = w - 1.05
    else:
        hx = x + .3
        hw = w - .55
    add_text(slide, hx, y + .18, hw, .38, heading, 15, NAVY, True)
    add_text(slide, x + .3, y + .72, w - .55, h - .88, body, body_size, MUTED)


def arrow(slide, x1, y1, x2, y2, color=BLUE, width=2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    c.line.end_arrowhead = True


def footer(slide, number):
    add_text(slide, 12.35, 7.05, .45, .25, str(number), 9, MUTED, align=PP_ALIGN.RIGHT)


# 1 — title
s = add_slide(dark=True)
add_shape(s, .78, 1.3, .1, 3.35, CYAN)
add_text(s, 1.2, 1.43, 11.3, 1.45,
         "AI-Assisted Red Teaming\nFramework", 38, WHITE, True)
add_text(s, 1.2, 3.1, 10.9, .95,
         "A policy-enforced, human-approved workflow for authorized penetration testing", 20,
         RGBColor(204, 224, 239))
add_shape(s, 1.2, 5.25, 10.8, .72, RGBColor(27, 52, 82), RGBColor(69, 112, 150), True)
add_text(s, 1.45, 5.43, 10.3, .35,
         "Engagement letter  →  Evidence-driven phases  →  Prioritized report", 15, WHITE, True,
         PP_ALIGN.CENTER)
add_text(s, 1.2, 6.62, 10.9, .3, "Project Phase II · Implementation", 12, RGBColor(165, 196, 219))

# 2 — real lifecycle
s = add_slide("The real penetration-testing lifecycle", "Six phases; authorization and evidence govern every transition")
phases = [
    ("Scope", "Written authorization\nTargets · ROE", NAVY),
    ("Recon", "Discover hosts, ports\nand technologies", BLUE),
    ("Analyze", "Correlate results\nand rank risk", CYAN),
    ("Exploit", "Minimal controlled\nverification", AMBER),
    ("Post-exploit", "Bounded impact\nproof only", RED),
    ("Report", "Evidence, priority\nand remediation", GREEN),
]
for i, (name, body, color) in enumerate(phases):
    x = .5 + i * 2.12
    add_shape(s, x, 1.72, 1.72, 1.72, color, radius=True)
    add_text(s, x + .12, 2.02, 1.48, .38, name, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, x + .12, 2.52, 1.48, .62, body, 10.5, WHITE, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + 1.75, 2.58, x + 2.08, 2.58, LINE, 1.5)
add_shape(s, .5, 3.88, 12.25, .65, PALE, LINE, True)
add_text(s, .75, 4.07, 11.75, .28,
         "Rules of engagement stay active across all phases: scope, permitted tools, timing and proof limits.",
         13, NAVY, True, PP_ALIGN.CENTER)
add_bullets(s, .8, 4.95, 11.7, 1.55, [
    "Each phase consumes evidence from the previous phase; exploitation is not a disconnected scan.",
    "A professional test proves claims while minimizing operational impact and preserving an audit trail.",
], 15)
footer(s, 2)

# 3 — system workflow mapping
s = add_slide("How the project satisfies every phase", "Standard phase → implemented system capability")
items = [
    ("Scoping & authorization", "Parse PDF/DOCX/text into targets, CIDRs, objectives, criticality, restrictions and exploit permission.", NAVY),
    ("Reconnaissance", "Draft nmap, DNS, HTTP, TLS and nuclei steps; approve and execute one command at a time.", BLUE),
    ("Vulnerability analysis", "Parse tool output, correlate and deduplicate findings, then calculate confidence, risk and priority.", CYAN),
    ("Exploitation", "Generate searchsploit, conservative sqlmap, declared-endpoint curl PoCs and Metasploit scanner steps from findings.", AMBER),
    ("Post-exploitation", "Only verified injection points; retrieve a DBMS banner or current DB user—no persistence or lateral movement.", RED),
    ("Reporting", "Render an evidence-backed HTML report with phase history, commands, verification status and remediation.", GREEN),
]
for i, (head, body, color) in enumerate(items):
    col, row = i % 2, i // 2
    card(s, .55 + col * 6.35, 1.25 + row * 1.78, 6.0, 1.48, head, body, color, i + 1, 11.7)
footer(s, 3)

# 4 — architecture
s = add_slide("System design", "Safety is enforced in the backend; the execution engine is replaceable")
add_shape(s, .45, 1.7, 1.45, .82, NAVY, radius=True)
add_text(s, .55, 1.88, 1.25, .4, "Operator", 14, WHITE, True, PP_ALIGN.CENTER)
add_shape(s, 2.3, 1.7, 1.7, .82, BLUE, radius=True)
add_text(s, 2.42, 1.88, 1.46, .4, "React UI", 14, WHITE, True, PP_ALIGN.CENTER)
add_shape(s, 4.4, 1.7, 2.0, .82, NAVY, radius=True)
add_text(s, 4.52, 1.88, 1.76, .4, "FastAPI API", 14, WHITE, True, PP_ALIGN.CENTER)
arrow(s, 1.92, 2.11, 2.25, 2.11)
arrow(s, 4.02, 2.11, 4.35, 2.11)
modules = [("Letter\nparser", NAVY), ("Planner", BLUE), ("Policy\nengine", RED),
           ("Executor", AMBER), ("Analyzer", CYAN), ("Reporter", GREEN)]
for i, (name, color) in enumerate(modules):
    x = .65 + i * 2.08
    add_shape(s, x, 3.05, 1.62, 1.02, PALE, color, True)
    add_text(s, x + .1, 3.28, 1.42, .55, name, 13, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    if i < 5:
        arrow(s, x + 1.65, 3.56, x + 2.03, 3.56, LINE, 1.5)
add_shape(s, 7.0, 1.43, 2.25, 1.36, WHITE, LINE, True)
add_text(s, 7.2, 1.67, 1.85, .35, "Persistence", 14, NAVY, True, PP_ALIGN.CENTER)
add_text(s, 7.2, 2.08, 1.85, .35, "SQLite · Fernet", 12, MUTED, align=PP_ALIGN.CENTER)
add_shape(s, 9.7, 1.43, 2.95, 1.36, WHITE, LINE, True)
add_text(s, 9.9, 1.67, 2.55, .35, "Execution engines", 14, NAVY, True, PP_ALIGN.CENTER)
add_text(s, 9.9, 2.08, 2.55, .4, "Docker/local  ·  Kali VM via SSH", 11.5, MUTED, align=PP_ALIGN.CENTER)
add_shape(s, .65, 4.75, 12.0, 1.25, PALE, LINE, True)
add_text(s, .9, 4.96, 2.2, .35, "Three mandatory gates", 16, NAVY, True)
add_text(s, 3.05, 4.92, 2.6, .65, "1  Letter permission\nExploit actions are explicit", 12.5, RED, True)
add_text(s, 6.0, 4.92, 2.7, .65, "2  Scope + flag policy\nUnknown actions fail closed", 12.5, BLUE, True)
add_text(s, 9.05, 4.92, 2.8, .65, "3  Human approval\nEvery command is reviewed", 12.5, GREEN, True)
add_text(s, .75, 6.35, 11.9, .35,
         "The plan is append-only by phase; executed evidence cannot be silently rewritten.", 13, MUTED,
         align=PP_ALIGN.CENTER)
footer(s, 4)

# 5 — safe execution
s = add_slide("Control plane and evidence flow", "A command becomes executable only after all checks pass")
steps = [
    ("Draft", "Planner proposes\na command", BLUE),
    ("Policy", "Tool, flags, scope\nand letter gate", RED),
    ("Approve", "Operator reviews\nexact command", GREEN),
    ("Execute", "Sandboxed process\nor Kali VM", AMBER),
    ("Analyze", "Parse output and\nlink evidence", CYAN),
]
for i, (head, body, color) in enumerate(steps):
    x = .65 + i * 2.48
    add_shape(s, x, 1.72, 1.9, 1.55, color, radius=True)
    add_text(s, x + .15, 1.98, 1.6, .35, head, 16, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, x + .15, 2.48, 1.6, .55, body, 11, WHITE, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + 1.94, 2.48, x + 2.42, 2.48, LINE, 1.5)
add_bullets(s, .75, 3.9, 5.8, 2.3, [
    "create_subprocess_exec: no command shell",
    "Timeouts, output caps and live terminal streaming",
    "VM runs record hostname, OS, kernel and SSH host-key fingerprint",
], 14, accent=BLUE)
add_bullets(s, 6.75, 3.9, 5.8, 2.3, [
    "Executed steps form an immutable audit trail",
    "Exploitation evidence upgrades scanner findings to verified",
    "Report preserves command, status, duration and evidence",
], 14, accent=GREEN)
footer(s, 5)

# 6 — tech stack
s = add_slide("Technology stack", "Small, inspectable components with deterministic offline fallbacks")
columns = [
    ("Frontend", ["React 19", "Vite", "Single-page control center"], BLUE),
    ("Backend", ["Python 3.10+", "FastAPI · Pydantic", "SQLAlchemy · SQLite", "Jinja2 reports"], NAVY),
    ("Security & I/O", ["Fernet encryption", "Paramiko SSH + tmux", "pypdf · python-docx"], GREEN),
    ("Assessment tools", ["nmap · dig · curl", "whatweb · sslscan · nuclei", "sqlmap · searchsploit", "msfconsole scanners"], AMBER),
]
for i, (head, values, color) in enumerate(columns):
    x = .55 + i * 3.17
    add_shape(s, x, 1.42, 2.82, 4.65, WHITE, color, True)
    add_shape(s, x, 1.42, 2.82, .75, color, radius=True)
    add_text(s, x + .15, 1.62, 2.52, .32, head, 15, WHITE, True, PP_ALIGN.CENTER)
    add_bullets(s, x + .22, 2.47, 2.38, 3.2, values, 12.5, accent=color)
add_text(s, .65, 6.42, 12.0, .35,
         "Deployment: Docker Compose for the toolkit/lab, or a real Kali attacker VM for demonstrable execution.",
         13, NAVY, True, align=PP_ALIGN.CENTER)
footer(s, 6)

# 7 — algorithms
s = add_slide("Core algorithms", "Deterministic logic keeps planning, correlation and safety explainable")
algos = [
    ("Scope validation", "IP ∈ CIDR · requested subnet ⊆ authorized subnet · domain suffix matching", RED),
    ("Allowlist policy", "Tokenize → identify tool → validate every flag/value → scope-check targets → fail closed", NAVY),
    ("Deduplication", "fingerprint = SHA-256(lower(title | endpoint | parameter))[:24]", BLUE),
    ("Evidence merge", "Worst severity + maximum score drivers + union of source tools + bounded evidence", CYAN),
    ("Findings-driven planning", "Version → searchsploit · URL → sqlmap · service → scanner module", AMBER),
    ("Verification matching", "Link recon and exploit findings when normalized endpoint or parameter keys overlap", GREEN),
]
for i, (head, body, color) in enumerate(algos):
    col, row = i % 2, i // 2
    card(s, .55 + col * 6.35, 1.25 + row * 1.8, 6.0, 1.5, head, body, color, body_size=12.1)
footer(s, 7)

# 8 — calculations
s = add_slide("Risk and priority calculations", "Risk describes the flaw; priority tells the client what to fix first")
add_shape(s, .7, 1.4, 5.75, 2.2, PALE, BLUE, True)
add_text(s, 1.0, 1.72, 5.15, .42, "Risk score", 19, NAVY, True, PP_ALIGN.CENTER)
add_text(s, 1.0, 2.3, 5.15, .55, "R = Exploitability × Impact × Exposure", 20, BLUE, True, PP_ALIGN.CENTER)
add_text(s, 1.0, 3.02, 5.15, .3, "Each driver: 1–5   →   range: 1–125", 12.5, MUTED, align=PP_ALIGN.CENTER)
add_shape(s, 6.88, 1.4, 5.75, 2.2, PALE, GREEN, True)
add_text(s, 7.18, 1.72, 5.15, .42, "Priority score", 19, NAVY, True, PP_ALIGN.CENTER)
add_text(s, 7.12, 2.2, 5.3, .9,
         "P = 0.40S + 0.25(E×20)\n    + 0.20A + 0.15C", 18, GREEN, True, PP_ALIGN.CENTER)
add_text(s, 7.18, 3.08, 5.15, .3, "All terms normalized to 0–100", 12.5, MUTED, align=PP_ALIGN.CENTER)
add_shape(s, .7, 4.08, 11.93, 1.72, WHITE, LINE, True)
add_text(s, 1.0, 4.34, 2.45, .4, "Inputs", 16, NAVY, True)
add_text(s, 3.05, 4.27, 8.95, .95,
         "S: severity anchor {Low 25, Medium 50, High 75, Critical 100}\n"
         "E: exploitability (1–5)   ·   A: asset criticality (0–100)   ·   C: evidence confidence (0–100)",
         13, INK)
add_text(s, 1.0, 5.25, 11.0, .3,
         "Example: High severity, E=4, A=80, C=90  →  P = 80;  with Impact=4 and Exposure=4  →  R = 64",
         13, NAVY, True, PP_ALIGN.CENTER)
add_text(s, .8, 6.3, 11.75, .35,
         "Drivers are clamped to legal ranges; missing values use safe defaults rather than zero.",
         12.5, MUTED, align=PP_ALIGN.CENTER)
footer(s, 8)

# 9 — project outcome
s = add_slide("What the project delivers", "Automation without surrendering authorization, restraint or accountability")
add_shape(s, .75, 1.35, 11.85, 1.0, NAVY, radius=True)
add_text(s, 1.0, 1.62, 11.35, .45,
         "Letter → scoped plan → approved execution → correlated findings → bounded verification → report",
         17, WHITE, True, PP_ALIGN.CENTER)
add_bullets(s, .95, 2.9, 5.7, 2.75, [
    "Full six-phase pentest lifecycle in one control center",
    "Offline deterministic operation; optional AI only proposes",
    "Real-system execution through Docker or a Kali VM",
], 15, accent=BLUE)
add_bullets(s, 6.75, 2.9, 5.7, 2.75, [
    "Scope, letter permissions and tool flags enforced in code",
    "Transparent scoring, deduplication and verification linking",
    "Client-ready report with reproducible evidence",
], 15, accent=GREEN)
add_shape(s, 1.6, 6.05, 10.1, .62, PALE, LINE, True)
add_text(s, 1.85, 6.2, 9.6, .3,
         "Core principle: automate repetitive toolwork; keep every consequential decision human-approved.",
         13.5, NAVY, True, PP_ALIGN.CENTER)
footer(s, 9)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
