"""Builds the project presentation (PPTX).

Run from the project root:
    .\\backend\\venv\\Scripts\\python.exe make_presentation.py

Produces Presentation.pptx: a plain-language walkthrough of (1) how a real
penetration test flows, (2) how this framework supports each phase, (3) the
parameters the framework takes and the basis on which each is defined or
calculated, and (4) the essential algorithms with commented code excerpts.

All algorithm excerpts are taken from the actual implementation in
backend/modules so the deck and the code cannot drift apart; comments are
trimmed for slide readability, never reinterpreted.
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Theme
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
LIGHT = RGBColor(0xDE, 0xEA, 0xF6)
INK = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x59, 0x59, 0x59)
CODEBG = RGBColor(0xF4, 0xF4, 0xF4)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xB4, 0x6A, 0x00)
RED = RGBColor(0xB3, 0x26, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)  # 16:9

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH


def slide():
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def box(s, x, y, w, h, fill=None, line=None, radius=False):
    from pptx.enum.shapes import MSO_SHAPE
    shape = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def text(s, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         font='Calibri', line_spacing=1.05, valign=None):
    """runs: a string, or a list of (text, dict-overrides) tuples for rich runs."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if valign:
        from pptx.enum.text import MSO_ANCHOR
        tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'mid': MSO_ANCHOR.MIDDLE,
                              'bottom': MSO_ANCHOR.BOTTOM}[valign]
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    first = True
    for para_runs in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = line_spacing
        for t, overrides in para_runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = overrides.get('font', font)
            f.size = Pt(overrides.get('size', size))
            f.bold = overrides.get('bold', bold)
            f.color.rgb = overrides.get('color', color)
    return tb


def title(s, t, sub=None):
    box(s, 0, 0, SW, Inches(1.0), fill=NAVY)
    text(s, Inches(0.55), Inches(0.18), SW - Inches(1.1), Inches(0.55), t,
         size=26, color=WHITE, bold=True)
    if sub:
        text(s, Inches(0.55), Inches(0.68), SW - Inches(1.1), Inches(0.3), sub,
             size=13, color=LIGHT)


def bullet_list(s, x, y, w, h, items, size=15, gap=6, color=INK, lead_color=None):
    rows = []
    for item in items:
        if isinstance(item, tuple):  # (lead, rest) — bold lead
            rows.append([(item[0], {'bold': True, 'color': lead_color or NAVY}),
                         (item[1], {'color': color})])
        else:
            rows.append([(item, {'color': color})])
    tb = text(s, x, y, w, h, rows, size=size, color=color)
    for p, _ in zip(tb.text_frame.paragraphs, items):
        p.space_after = Pt(gap)
    return tb


def code(s, x, y, w, h, code_text, size=11):
    box(s, x, y, w, h, fill=CODEBG, line=RGBColor(0xC9, 0xC9, 0xC9))
    tb = s.shapes.add_textbox(x + Inches(0.12), y + Inches(0.08),
                              w - Inches(0.24), h - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = False
    for i, line in enumerate(code_text.strip('\n').split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.name = 'Consolas'
        r.font.size = Pt(size)
        # comments in green, the rest in ink
        stripped = line.lstrip()
        if stripped.startswith('#'):
            r.font.color.rgb = GREEN
        else:
            r.font.color.rgb = INK
        p.line_spacing = 1.0
    return tb


def arrows(s, points, color=BLUE, weight=2.5):
    from pptx.enum.shapes import MSO_CONNECTOR
    for (x1, y1, x2, y2) in points:
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        c.line.color.rgb = color
        c.line.width = Pt(weight)
        c.shadow.inherit = False


def chip(s, x, y, w, h, label, fill, tcolor=WHITE, size=13):
    box(s, x, y, w, h, fill=fill, radius=True)
    text(s, x, y, w, h, label, size=size, color=tcolor, bold=True,
         align=PP_ALIGN.CENTER, valign='mid')


# ---------------------------------------------------------------------------
# 1. Title
s = slide()
box(s, 0, 0, SW, SH, fill=NAVY)
box(s, Inches(0.9), Inches(1.7), Inches(0.14), Inches(2.6), fill=BLUE)
text(s, Inches(1.3), Inches(1.8), Inches(11), Inches(1.6),
     'A Policy-Enforced Framework for\nSemi-Automated Penetration Testing',
     size=40, color=WHITE, bold=True, line_spacing=1.05)
text(s, Inches(1.3), Inches(3.5), Inches(11), Inches(0.6),
     'From an authorized engagement letter to verified, prioritized and reported findings — '
     'with every command policy-checked and human-approved.',
     size=17, color=LIGHT)
text(s, Inches(1.3), Inches(5.9), Inches(11), Inches(0.9),
     [[('Project Phase — II  ·  Implementation', {'size': 15, 'color': LIGHT})],
      [('Backend · Python (FastAPI)   |   Frontend · React   |   Tooling · Kali / Docker lab', {'size': 13, 'color': RGBColor(0x9D, 0xC3, 0xE6)})]])

# ---------------------------------------------------------------------------
# 2. Agenda
s = slide()
title(s, 'What this presentation covers', 'Plain language, one idea per slide')
items = [
    ('Part 1 — the real-world flow:', ' how a professional penetration test is actually conducted, phase by phase.'),
    ('Part 2 — the framework:', ' what the system does at each of those phases, and its overall architecture.'),
    ('Part 3 — the parameters:', ' every input the framework takes, and the basis on which each is defined or calculated.'),
    ('Part 4 — the algorithms:', ' the essential logic — scoring, deduplication, safe planning, policy validation, verification matching — shown as commented code from the actual implementation.'),
]
bullet_list(s, Inches(0.8), Inches(1.8), Inches(11.8), Inches(4.2), items, size=18, gap=18)
text(s, Inches(0.8), Inches(6.3), Inches(11.8), Inches(0.6),
     'A live demo of the complete flow — letter import to report — accompanies this deck.',
     size=14, color=MUTED)

# ---------------------------------------------------------------------------
# 3. The standard penetration-testing lifecycle (diagram)
s = slide()
title(s, 'How a real penetration test flows', 'The standard lifecycle every engagement follows (PTES / OWASP guidance)')
phases = [
    ('1  Scoping &\nAuthorization', NAVY),
    ('2  Recon &\nEnumeration', BLUE),
    ('3  Vulnerability\nAnalysis', BLUE),
    ('4  Exploitation', AMBER),
    ('5  Post-\nExploitation', AMBER),
    ('6  Reporting', GREEN),
]
w, gap = Inches(1.85), Inches(0.28)
x0 = Inches(0.55)
for i, (label, color) in enumerate(phases):
    x = x0 + i * (w + gap)
    box(s, x, Inches(2.0), w, Inches(1.25), fill=color, radius=True)
    text(s, x, Inches(2.0), w, Inches(1.25), label, size=13.5, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, valign='mid')
    if i < 5:
        arrows(s, [(x + w + Inches(0.03), Inches(2.62), x + w + gap - Inches(0.03), Inches(2.62))])
# rules-of-engagement bar under all phases
box(s, x0, Inches(3.55), 6 * w + 5 * gap, Inches(0.5), fill=LIGHT, radius=True)
text(s, x0, Inches(3.55), 6 * w + 5 * gap, Inches(0.5),
     'Authorization letter governs the whole engagement — scope, allowed tools, what may be exploited',
     size=12.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign='mid')
bullet_list(s, Inches(0.55), Inches(4.35), Inches(12.2), Inches(2.7), [
    ('Two golden constraints hold across every phase: ', 'stay inside the authorized scope, and keep evidence for everything you claim.'),
    ('Phases 1 and 6 are human work today — ', 'law and liability, respectively. Phases 2–5 are repetitive, tool-driven and evidence-heavy.'),
    ('The feedback loop is what makes it slow: ', 'what you find in phase 3 decides what you do in phase 4; humans currently carry that hand-off manually.'),
], size=14.5, gap=10)

# ---------------------------------------------------------------------------
# 4. What each phase involves
s = slide()
title(s, 'What actually happens in each phase', 'The work a tester does — and where the hours go')
rows = [
    ('Phase', 'The tester’s actual work', 'Where the time goes'),
    ('1  Scoping', 'Read the engagement letter; extract targets, scope, criticality, tool restrictions', 'Manual reading; easy to misread scope'),
    ('2  Recon', 'Run nmap, dig, curl, whatweb, sslscan, nuclei… on each target', 'Recalling the right commands and flags'),
    ('3  Analysis', 'Read every output; correlate across tools; dedupe; judge severity', 'The real bottleneck — cross-referencing'),
    ('4  Exploitation', 'Verify each suspected flaw with a careful, minimal proof-of-concept', 'Crafting safe PoCs; staying in scope'),
    ('5  Post-exploitation', 'Prove impact only as far as the letter allows; never more', 'Restraint — proving without damaging'),
    ('6  Reporting', 'Write the cited, prioritized report the client will act on', '1–3 hours of writing per engagement'),
]
y = Inches(1.35)
col_x = [Inches(0.55), Inches(2.6), Inches(7.6)]
col_w = [Inches(2.0), Inches(4.95), Inches(5.1)]
for i, row in enumerate(rows):
    h = Inches(0.82)
    hdr = i == 0
    for j, cell in enumerate(row):
        box(s, col_x[j], y, col_w[j], h - Inches(0.04),
            fill=NAVY if hdr else (LIGHT if i % 2 else WHITE),
            line=RGBColor(0xC9, 0xD8, 0xEA))
        text(s, col_x[j] + Inches(0.1), y, col_w[j] - Inches(0.2), h - Inches(0.04),
             cell, size=12.5, bold=hdr, color=WHITE if hdr else INK, valign='mid')
    y += h

# ---------------------------------------------------------------------------
# 5. The idea
s = slide()
title(s, 'The design idea', 'Automate the toolwork, keep the human in command')
box(s, Inches(0.7), Inches(1.5), Inches(5.9), Inches(4.6), fill=LIGHT, radius=True)
text(s, Inches(1.0), Inches(1.7), Inches(5.3), Inches(0.5), 'What the system does', size=18, color=NAVY, bold=True)
bullet_list(s, Inches(1.0), Inches(2.25), Inches(5.3), Inches(3.6), [
    'Reads the authorization letter itself (PDF) and builds the scope',
    'Drafts every command a phase needs — policy-checked before you see it',
    'Executes on your approval only, recording full evidence',
    'Correlates, dedupes and scores every finding automatically',
    'Proposes the next phase’s steps from what was actually found',
    'Writes the cited, prioritized report',
], size=14, gap=8)
box(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(4.6), fill=WHITE, line=BLUE, radius=True)
text(s, Inches(7.15), Inches(1.7), Inches(5.3), Inches(0.5), 'What stays with the human', size=18, color=BLUE, bold=True)
bullet_list(s, Inches(7.15), Inches(2.25), Inches(5.3), Inches(3.6), [
    'Every command is approved command-by-command (human-in-the-loop)',
    'Scope, restrictions and exploitation rights are enforced in code — not left to vigilance',
    'Exploit-grade actions are gated by the letter itself',
    'Post-exploitation is bounded to a single proof record per flaw',
    'The operator can veto, edit or disable any step before it runs',
], size=14, gap=8)
text(s, Inches(0.7), Inches(6.35), Inches(12), Inches(0.5),
     'Principle: fail closed — anything not explicitly allowed by policy and the letter is refused, with the reason shown.',
     size=14, color=RED, bold=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 6. Architecture
s = slide()
title(s, 'System architecture', 'How the pieces fit together')
# operator
chip(s, Inches(0.5), Inches(1.75), Inches(1.7), Inches(0.75), 'Operator\n(browser UI)', NAVY, size=12)
# frontend
box(s, Inches(2.6), Inches(1.75), Inches(1.9), Inches(0.75), fill=BLUE, radius=True)
text(s, Inches(2.6), Inches(1.75), Inches(1.9), Inches(0.75), 'React frontend', size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign='mid')
# API
box(s, Inches(4.9), Inches(1.75), Inches(1.9), Inches(0.75), fill=NAVY, radius=True)
text(s, Inches(4.9), Inches(1.75), Inches(1.9), Inches(0.75), 'FastAPI backend', size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign='mid')
arrows(s, [(Inches(2.2), Inches(2.12), Inches(2.6), Inches(2.12)),
           (Inches(4.5), Inches(2.12), Inches(4.9), Inches(2.12))])
# module row
mods = [('Letter\nparser',), ('Plan\nagent',), ('Policy\nengine',), ('Executor\n(local/VM)',), ('Analyzer\nagent',), ('Report\ngenerator',)]
mw, mgap = Inches(1.35), Inches(0.22)
mx0 = Inches(3.1)
for i, (label,) in enumerate(mods):
    x = mx0 + i * (mw + mgap)
    box(s, x, Inches(3.1), mw, Inches(1.0), fill=LIGHT, line=BLUE, radius=True)
    text(s, x, Inches(3.1), mw, Inches(1.0), label, size=11.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign='mid')
    arrows(s, [(x + mw // 2, Inches(3.05), x + mw // 2, Inches(3.1))], color=RGBColor(0x9D, 0xC3, 0xE6), weight=1.5)
# database
box(s, Inches(3.1), Inches(4.55), Inches(1.5), Inches(0.6), fill=WHITE, line=NAVY, radius=True)
text(s, Inches(3.1), Inches(4.55), Inches(1.5), Inches(0.6), 'SQLite DB', size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER, valign='mid')
arrows(s, [(Inches(5.85), Inches(4.1), Inches(4.6), Inches(4.7))], color=RGBColor(0x9D, 0xC3, 0xE6), weight=1.5)
# execution targets
box(s, Inches(8.6), Inches(3.1), Inches(4.2), Inches(1.0), fill=AMBER, radius=True)
text(s, Inches(8.7), Inches(3.1), Inches(4.0), Inches(1.0),
     'Execution engines\nDocker (lab)  ·  Kali VM over SSH\n(each run carries an attestation)', size=12.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign='mid')
arrows(s, [(Inches(7.35), Inches(3.6), Inches(8.6), Inches(3.6))])
# letter + report artifacts
box(s, Inches(0.5), Inches(3.1), Inches(1.7), Inches(1.0), fill=WHITE, line=MUTED, radius=True)
text(s, Inches(0.5), Inches(3.1), Inches(1.7), Inches(1.0), 'Engagement\nletter (PDF)', size=12, color=INK, bold=True, align=PP_ALIGN.CENTER, valign='mid')
arrows(s, [(Inches(2.2), Inches(3.6), Inches(3.1), Inches(3.6))], color=MUTED, weight=1.5)
box(s, Inches(10.9), Inches(5.0), Inches(1.9), Inches(0.7), fill=GREEN, radius=True)
text(s, Inches(10.9), Inches(5.0), Inches(1.9), Inches(0.7), 'HTML report', size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign='mid')
arrows(s, [(Inches(11.85), Inches(4.75), Inches(11.85), Inches(5.0))])
bullet_list(s, Inches(0.55), Inches(5.8), Inches(12.2), Inches(1.5), [
    ('Every module is replaceable: ', 'an AI provider can draft plans and analyze outputs, but the deterministic path works fully offline.'),
    ('Safety lives in the backend, not the UI: ', 'commands are re-validated at plan time, at draft time, and again at approval time.'),
], size=14, gap=8)

# ---------------------------------------------------------------------------
# 7. Phase mapping
s = slide()
title(s, 'Framework support for each lifecycle phase', 'Standard pentest phase → what the system contributes')
rows = [
    ('Lifecycle phase', 'Framework capability', 'Where it lives'),
    ('1  Scoping', 'Letter PDF parsed into structured scope: targets, CIDRs, criticality, per-target tool restrictions, verification endpoints, exploitation authorization', 'engagement_parser.py'),
    ('2  Recon', 'Deterministic or AI-drafted plan per target: nmap, traceroute, dig, curl, whatweb, sslscan, nuclei — all rate-limited and scope-checked', 'planner.py'),
    ('3  Analysis', 'Per-tool output parsers correlate and dedupe raw outputs into scored findings (severity, risk, priority, confidence)', 'analyzer.py'),
    ('4  Exploitation', 'Findings-driven plan: searchsploit lookups per version fingerprint, sqlmap verification, letter-declared PoC endpoints, msfconsole scanners', 'exploit_planner.py'),
    ('5  Post-exploitation', 'Bounded to a single proof record per verified flaw (DBMS banner, current user) — nothing more, ever', 'exploit_planner.py'),
    ('6  Reporting', 'Cited HTML report: engagement reference, phase timeline, verified evidence, every executed command with its result', 'reporter.py'),
]
y = Inches(1.35)
col_x = [Inches(0.55), Inches(2.5), Inches(9.1)]
col_w = [Inches(1.9), Inches(6.55), Inches(3.6)]
for i, row in enumerate(rows):
    h = Inches(0.79)
    hdr = i == 0
    for j, cell in enumerate(row):
        box(s, col_x[j], y, col_w[j], h - Inches(0.04),
            fill=NAVY if hdr else (LIGHT if i % 2 else WHITE),
            line=RGBColor(0xC9, 0xD8, 0xEA))
        text(s, col_x[j] + Inches(0.1), y, col_w[j] - Inches(0.2), h - Inches(0.04),
             cell, size=11.5, bold=hdr, color=WHITE if hdr else INK, valign='mid',
             font='Consolas' if (j == 2 and not hdr) else 'Calibri')
    y += h
text(s, Inches(0.55), Inches(6.75), Inches(12.2), Inches(0.55),
     [[('Throughout every phase: ', {'bold': True, 'color': NAVY}),
       ('allowlist policy engine + human-in-the-loop approval + full audit trail  (policy_engine.py, executor.py)', {})]],
     size=13)

# ---------------------------------------------------------------------------
# 8. The phase model and its gates (diagram)
s = slide()
title(s, 'The engagement state machine', 'Each phase must be analyzed before the next can be drafted')
flow = [
    ('Letter\nimport', NAVY),
    ('Recon plan\n+ execute', BLUE),
    ('Analyze\nfindings', BLUE),
    ('Exploitation plan\n+ execute', AMBER),
    ('Analyze\nverification', AMBER),
    ('Post-exploit\n(bounded)', AMBER),
    ('Report', GREEN),
]
w, gap = Inches(1.55), Inches(0.3)
x0 = Inches(0.55)
for i, (label, color) in enumerate(flow):
    x = x0 + i * (w + gap)
    box(s, x, Inches(1.7), w, Inches(1.05), fill=color, radius=True)
    text(s, x, Inches(1.7), w, Inches(1.05), label, size=11.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER, valign='mid')
    if i < len(flow) - 1:
        arrows(s, [(x + w + Inches(0.04), Inches(2.22), x + w + gap - Inches(0.04), Inches(2.22))])
# gates
gates = [
    ('Gate 1 — letter gate', 'Exploitation-grade commands (sqlmap, msfconsole, curl with a body) are refused unless the letter authorizes controlled exploitation for that target.', Inches(3.3)),
    ('Gate 2 — scope gate', 'Every command’s target address must fall inside an authorized scope (IP, CIDR or domain); resolvers must be public or in-scope.', Inches(4.45)),
    ('Gate 3 — flag gate', 'Only enumerated flags per tool are legal; file-writes, redirects, listeners, shells and data dumps have no legal spelling.', Inches(5.6)),
]
for label, desc, gy in gates:
    chip(s, Inches(0.55), gy, Inches(2.3), Inches(0.85), label.split(' — ')[0] + '\n' + label.split(' — ')[1], RED, size=11)
    text(s, Inches(3.05), gy + Inches(0.05), Inches(9.7), Inches(0.75), desc, size=12.5, valign='mid')
text(s, Inches(0.55), Inches(6.7), Inches(12.2), Inches(0.5),
     'All three gates run again at every approval — a plan that passed yesterday’s policy can never ride through today’s review unexamined.',
     size=13.5, color=NAVY, bold=True)

# ---------------------------------------------------------------------------
# PART 3 — parameters
# 9. Scoping parameters (from the letter)
s = slide()
title(s, 'Input parameters — what the framework takes', 'Part 3: every parameter, and the basis on which it is defined')
rows = [
    ('Parameter', 'Taken from', 'Basis / rationale'),
    ('Target address', 'Engagement letter (Section 3)', 'The letter names each authorized host/CIDR; only these may appear in any command. Host-only tools (nmap, dig) get the bare host, web tools get host:port.'),
    ('Authorized scopes', 'Letter scope statements', 'A list of IPs, CIDRs and domains. The policy engine checks every command target against this list — a subnet must fit inside an authorized network, not just its base address.'),
    ('Asset criticality', 'Letter criticality ratings', '0–100, client-declared business criticality. Default 70 when the letter does not say. Feeds the priority score of every finding on that target.'),
    ('Restricted tools', 'Letter rules of engagement', 'Per-target deny-list parsed from the letter; restricted steps are dropped at plan time and refused at approval time.'),
    ('Verification endpoints', 'Letter verification section', 'Exact URLs the client authorizes for proof-of-concept testing (e.g. a login endpoint). The only places a PoC payload may be sent.'),
    ('Exploitation flag', 'Letter authorization clause', 'Boolean: does the letter permit controlled verification? Gates the entire exploitation and post-exploitation phase for that target.'),
    ('Objective / prompt', 'Operator', 'Free text steering plan drafting — “audit security headers” shapes the drafted commands when an AI provider is configured.'),
]
y = Inches(1.35)
col_x = [Inches(0.55), Inches(2.75), Inches(5.35)]
col_w = [Inches(2.15), Inches(2.55), Inches(7.35)]
for i, row in enumerate(rows):
    h = Inches(0.71)
    hdr = i == 0
    for j, cell in enumerate(row):
        box(s, col_x[j], y, col_w[j], h - Inches(0.04),
            fill=NAVY if hdr else (LIGHT if i % 2 else WHITE),
            line=RGBColor(0xC9, 0xD8, 0xEA))
        text(s, col_x[j] + Inches(0.1), y, col_w[j] - Inches(0.2), h - Inches(0.04),
             cell, size=10.5, bold=hdr, color=WHITE if hdr else INK, valign='mid')
    y += h

# ---------------------------------------------------------------------------
# 10. Scoring parameters
s = slide()
title(s, 'Input parameters — the finding-scoring drivers', 'Each driver, its range, and where the value comes from')
rows = [
    ('Driver', 'Range', 'Defined by / calculated from'),
    ('severity', 'Low · Medium · High · Critical', 'The tool that produced the finding (nuclei level, sslscan protocol, sqlmap = Critical). Normalized to the four-level scale; AI mode maps to the same scale.'),
    ('exploitability', '1–5', 'How readily the issue is weaponized by a competent attacker — set per finding-type in the analyzer (e.g. confirmed SQLi = 5, banner disclosure = 2).'),
    ('impact', '1–5', 'Damage if exploited — e.g. session-token theft = 4, deprecated TLS = 4, referrer leak = 2.'),
    ('exposure', '1–5', 'Reachability: remotely reachable from the network = 4–5, requires adjacent access = 2. nmap "open" = 4, "filtered" = 2.'),
    ('confidence_score', '0–100', 'Evidence strength: version-confirmed nmap banner = 95, tentative service guess = 55, template signature match = 80, confirmed injection = 95.'),
    ('asset_criticality', '0–100', 'Client-declared (from the letter); falls back to this target’s criticality, then the global default of 70.'),
]
y = Inches(1.35)
col_x = [Inches(0.55), Inches(2.6), Inches(5.5)]
col_w = [Inches(2.0), Inches(2.85), Inches(7.2)]
for i, row in enumerate(rows):
    h = Inches(0.83)
    hdr = i == 0
    for j, cell in enumerate(row):
        box(s, col_x[j], y, col_w[j], h - Inches(0.04),
            fill=NAVY if hdr else (LIGHT if i % 2 else WHITE),
            line=RGBColor(0xC9, 0xD8, 0xEA))
        text(s, col_x[j] + Inches(0.1), y, col_w[j] - Inches(0.2), h - Inches(0.04),
             cell, size=11, bold=hdr, color=WHITE if hdr else INK, valign='mid')
    y += h
text(s, Inches(0.55), Inches(6.6), Inches(12.2), Inches(0.5),
     'All drivers are bounded integers — a missing or malformed value falls back to a safe default, never to an error or a zero.',
     size=13.5, color=NAVY, bold=True)

# ---------------------------------------------------------------------------
# PART 4 — algorithms
# 11. Priority + risk
s = slide()
title(s, 'Essential algorithms — 1. Finding score', 'Priority (fix-first ordering) and risk (raw risk product) — from analyzer.py')
code(s, Inches(0.55), Inches(1.4), Inches(7.1), Inches(3.35), '''
# Severity anchor: each level maps to a fixed scale value.
SEVERITY = {'Low': 25, 'Medium': 50, 'High': 75, 'Critical': 100}

def score_finding(finding):
    severity = normalize_severity(finding.get('severity'))
    severity_score = SEVERITY[severity]              # Low..Critical -> 25..100
    # Each driver is clamped into its legal range, with a
    # default for missing/malformed values (fail safe, not zero):
    exploitability = bounded_int(finding.get('exploitability'), 3, 1, 5)
    impact        = bounded_int(finding.get('impact'),        3, 1, 5)
    exposure      = bounded_int(finding.get('exposure'),      3, 1, 5)
    confidence    = bounded_int(finding.get('confidence_score'),  70, 0, 100)
    asset_crit    = bounded_int(finding.get('asset_criticality'), 70, 0, 100)

    # RISK = how dangerous the flaw itself is:
    # three 1-5 drivers multiplied -> 1..125.
    risk = exploitability * impact * exposure

    # PRIORITY = what the client should fix FIRST:
    # 40% severity + 25% exploitability + 20% asset criticality
    #          + 15% confidence  (each term scaled to 100 -> 0..100)
    priority = round(.40 * severity_score
                   + .25 * exploitability * 20
                   + .20 * asset_crit
                   + .15 * confidence)
    return severity, risk, priority, confidence
''', size=10.5)
bullet_list(s, Inches(8.0), Inches(1.5), Inches(4.8), Inches(4.6), [
    ('Why a weighted sum for priority? ', 'Severity alone ignores context: a High on a business-critical asset outranks a Critical on a lab box. The weights encode remediation order, not danger.'),
    ('Why a product for risk? ', 'Exploitability, impact and exposure multiply: any one near zero makes the flaw negligible — a weighted sum would not do that.'),
    ('Why the 40/25/20/15 split? ', 'Severity dominates (it is the assessed level); exploitability next (fix what is weaponizable first); the asset’s business weight; then evidence confidence.'),
], size=13, gap=12)

# ---------------------------------------------------------------------------
# 12. Dedupe fingerprint + merge
s = slide()
title(s, 'Essential algorithms — 2. Cross-tool deduplication', 'Different tools reporting the same flaw become one finding — from analyzer.py')
code(s, Inches(0.55), Inches(1.4), Inches(7.1), Inches(2.75), '''
def fingerprint(finding):
    # Identity of a flaw = (title, endpoint, parameter), lowercased.
    # nmap "port 3000 open" and curl "missing CSP" on the same host
    # differ in title -> two findings; two tools reporting the same
    # missing header -> identical material -> ONE finding.
    material = '|'.join(str(finding.get(k, '')).lower().strip()
                        for k in ('title', 'endpoint', 'parameter'))
    # SHA-256 truncated to 24 hex chars: a stable, collision-resistant key
    # that also works as a database identity across analysis runs.
    return hashlib.sha256(material.encode()).hexdigest()[:24]
''', size=10.5)
code(s, Inches(0.55), Inches(4.35), Inches(7.1), Inches(2.5), '''
def _combine(current, incoming):   # fold a duplicate into the held finding
    # Keep the WORST severity observed, never the first-seen one:
    if SEVERITY_RANK[incoming['severity']] > SEVERITY_RANK[current['severity']]:
        current['severity'] = incoming['severity']
    # Every scoring driver keeps its strongest observed value...
    for field, default in SCORE_DRIVERS:
        current[field] = max(current[field], incoming[field])
    # ...tools are unioned (evidence the flaw is multi-tool)...
    current['source_tools'] = set(current['source_tools']) | set(incoming['source_tools'])
    # ...and evidence lines are appended (bounded), so proof accumulates.
''', size=10.5)
bullet_list(s, Inches(8.0), Inches(1.5), Inches(4.8), Inches(5.3), [
    ('The merge is order-independent: ', 'which tool ran first cannot decide a finding’s severity — the maximum wins.'),
    ('Then the merged finding is rescored, ', 'so severity, risk and priority always agree with the strongest evidence actually held.'),
    ('Scale guard: ', 'a 50-step plan cannot inflate one finding — evidence text is capped at 20,000 chars per finding and 200 findings per analysis.'),
], size=13, gap=12)

# ---------------------------------------------------------------------------
# 13. Policy engine
s = slide()
title(s, 'Essential algorithms — 3. Command policy validation', 'Allowlist, fail-closed — every command passes this before any human sees it — from policy_engine.py')
code(s, Inches(0.55), Inches(1.35), Inches(7.6), Inches(5.35), '''
def validate_command(command, authorized_scopes,
                     expected_tool=None, allow_exploitation=False):
    # 1. Shell-split the command (it later runs via exec, not a shell,
    #    so metacharacters are inert; only control chars are refused).
    tokens = shlex.split(command)
    tool = tokens[0]

    # 2. THE EXPLOITATION GATE: sqlmap/msfconsole, or curl carrying a
    #    request body (a PoC payload), are exploitation-grade. Without
    #    the letter's authorization the command is refused HERE.
    if exploitation_grade and not allow_exploitation:
        return False, 'letter does not authorize exploitation', rules

    # 3. THE FLAG WALK: walk the argument list against this tool's
    #    enumerated spec (bool flags, value flags, target flags...).
    #    Anything not enumerated -> refused. -o (write file),
    #    --os-shell, -x/--proxy, --dump have no legal spelling.
    targets, resolvers, rejected = self.scan_arguments(tokens, rules)
    if rejected is not None:
        return False, f'Blocked flag: {rejected}', rules

    # 4. THE SCOPE CHECK: every target must validate against the
    #    authorized scopes (IP-in-CIDR, domain suffix, subnet_of for
    #    ranges). Resolvers must be public (8.8.8.8...) or in scope.
    if any(not self.validate_target(t, authorized_scopes) for t in targets):
        return False, 'target outside authorized scope', rules

    # 5. Only now is the command reviewable — HITL approval still
    #    stands between this and execution.
    return True, 'allowed; HITL approval required', rules
''', size=10.5)
bullet_list(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.2), [
    ('Allowlist, not blocklist: ', 'a forgotten dangerous flag fails closed instead of slipping through the gaps.'),
    ('Per-tool flag semantics: ', '“nmap -A” takes no value, “curl -A” does — one shared table cannot express both.'),
    ('Refusals carry the reason: ', 'the operator sees the letter’s decision, not a silent error.'),
    ('msfconsole scripts get their own ', 'statement-by-statement validation: only use/set/run; scanner/exploit module trees; RHOSTS scope-checked.'),
], size=12.5, gap=10)

# ---------------------------------------------------------------------------
# 14. Findings-driven exploitation planning
s = slide()
title(s, 'Essential algorithms — 4. Findings-driven phase planning', 'The next phase’s plan is drafted FROM what was actually found — from exploit_planner.py')
code(s, Inches(0.55), Inches(1.35), Inches(7.6), Inches(5.35), '''
def draft_exploitation_plan(findings, target, scopes,
                            verification_endpoints=()):
    # INPUT: the scored findings of the analysis phase.
    # OUTPUT: policy-validated verification steps.

    # (a) Version fingerprints -> offline Exploit-DB lookups.
    #     nmap titles carry "Exposed ssh service on port 22 (OpenSSH 4.7)"
    #     -> searchsploit ssh 4.7      (capped at 5 lookups)
    for service, version in _version_fingerprints(findings):
        steps.append('searchsploit', f'searchsploit {service} {version}')

    # (b) Web endpoints found by recon -> conservative sqlmap
    #     verification: --batch (never asks), risk/level 1 (shallow),
    #     technique B (boolean-based, no time-delays).  (cap: 3)
    for url in _web_endpoints(findings, target):
        steps.append('sqlmap', f'sqlmap -u {url} {SQLMAP_VERIFY_FLAGS}')

    # (c) The letter's OWN declared endpoints -> the exact PoC the
    #     client authorized, e.g. an auth-bypass login payload.
    for path in verification_endpoints:
        steps.append('curl', f'curl --data "{SQLI_LOGIN_PAYLOAD}" {base}{path}')

    # (d) Exposed services -> msfconsole SCANNER modules only
    #     (banner readers, not payloads), run on the attacker VM.
    steps.extend(_msf_scanner_steps(findings, scopes))

    # Every step is policy-validated BEFORE it is returned:
    # a planner bug yields a useless step, never an unchecked command.
''', size=10.5)
bullet_list(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(5.2), [
    ('Plans are data, not promises: ', 'each drafted step is re-run through the full policy engine before the operator ever sees it.'),
    ('Caps everywhere: ', 'max 5 searchsploit lookups, 3 sqlmap runs — a fingerprint-rich target cannot flood the phase or the timeout.'),
    ('Post-exploitation is stricter still: ', 'only sqlmap-VERIFIED injection points qualify, and only for two one-line facts — DBMS banner and current user.'),
], size=12.5, gap=10)

# ---------------------------------------------------------------------------
# 15. Verification matching
s = slide()
title(s, 'Essential algorithms — 5. Evidence verification matching', 'Linking each exploit result back to the recon finding it proves — from analyzer.py')
code(s, Inches(0.55), Inches(1.4), Inches(7.6), Inches(4.9), '''
def match_verification(recon_findings, exploitation_findings):
    # Build identity keys for one finding:
    #   - its endpoint, normalized ('host:port/tcp' and
    #     'http://host:port' reduce to the same bare 'host:port')
    #   - its parameter name, if any (sqlmap reports 'email')
    def keys(finding):
        found, endpoint = set(), (finding.get('endpoint') or '').strip()
        if endpoint:
            found.add(endpoint)
            bare = re.sub(r'^[a-z]+://', '', endpoint)
            bare = bare.split('/tcp')[0].split('/')[0]
            found.add(bare.lower())
        if finding.get('parameter'):
            found.add('param:' + finding['parameter'].lower())
        return found

    # A recon finding and an exploitation finding MATCH when
    # their key sets overlap: same host:port, same URL, or same
    # injectable parameter.
    for recon in recon_findings:
        for exploit in exploitation_findings:
            if keys(recon) & keys(exploit):
                links.append((recon, exploit, exploit['evidence']))
    # The matched recon finding is upgraded to 'verified',
    # carrying the exploit's evidence as its proof.
''', size=10.5)
bullet_list(s, Inches(8.5), Inches(1.5), Inches(4.3), Inches(4.9), [
    ('Why loose on purpose: ', 'a false link only strengthens one finding’s evidence attribution — it can never authorize a command.'),
    ('The report distinguishes ', '“signature says so” (scanner evidence) from “exploitation confirmed” (verified) — the client prioritizes differently.'),
    ('Two independent sources agreeing ', '(nmap fingerprint + Metasploit scanner banner) raise confidence in every version-mapped finding.'),
], size=12.5, gap=10)

# ---------------------------------------------------------------------------
# 16. Safe execution + reporting
s = slide()
title(s, 'Execution safety and the report', 'What happens between approval and evidence')
bullet_list(s, Inches(0.55), Inches(1.4), Inches(6.3), Inches(5.2), [
    ('Human-in-the-loop: ', 'nothing executes without a per-command approval; the full command, not just its description, is what the operator approves.'),
    ('Two execution engines: ', 'Docker lab mode (local containers) or attacker-VM mode — commands are typed into a tmux session on the Kali VM over SSH, visible on the VM’s own console.'),
    ('Per-run attestation: ', 'each VM-mode step records the machine’s hostname, OS, kernel and SSH host-key fingerprint — verifiable proof of where every command ran.'),
    ('Hard bounds: ', '360-second timeout per command with interrupt; output caps; live streaming terminal for the operator.'),
    ('The report: ', 'cites the engagement reference, walks the phase timeline, separates verified from signature findings, and lists every command with exit code, duration and evidence — the reproducibility record.'),
], size=13.5, gap=10)
box(s, Inches(7.1), Inches(1.5), Inches(5.7), Inches(5.0), fill=LIGHT, radius=True)
text(s, Inches(7.4), Inches(1.7), Inches(5.1), Inches(0.5), 'The safety ledger', size=17, color=NAVY, bold=True)
bullet_list(s, Inches(7.4), Inches(2.3), Inches(5.1), Inches(4.1), [
    'Allowlist policy engine — fail closed',
    'Letter-gated exploitation (per target)',
    'Bounded post-exploitation (one record per flaw)',
    'Scope-checked targets, resolvers, msf RHOSTS',
    'Per-command HITL approval',
    'Full audit trail + per-run attestation',
    'Encrypted secret storage (Fernet) for VM credentials',
], size=13, gap=8)

# ---------------------------------------------------------------------------
# 17. Summary
s = slide()
box(s, 0, 0, SW, SH, fill=NAVY)
text(s, Inches(0.9), Inches(0.9), Inches(11.5), Inches(0.8),
     'In summary', size=34, color=WHITE, bold=True)
bullet_list(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.2), [
    ('The lifecycle is respected, not replaced: ', 'every standard phase — scoping, recon, analysis, exploitation, post-exploitation, reporting — maps to a framework capability.'),
    ('The parameters are principled: ', 'scope comes from the client’s letter; scoring drivers are bounded, evidence-derived values feeding a weighted priority and multiplicative risk.'),
    ('The algorithms are deterministic and inspectable: ', 'deduplication by content fingerprint, phase planning from actual findings, policy validation by allowlist, verification by endpoint/parameter matching.'),
    ('The human stays in command: ', 'automation drafts and executes only what policy, the letter, and a per-command approval allow — and the report proves it.'),
], size=17, gap=16, color=WHITE, lead_color=RGBColor(0x9D, 0xC3, 0xE6))
text(s, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.5),
     'Demonstrated from a Kali attacker VM against a real self-hosted application (OmniRoute), driven end-to-end from the engagement letter.',
     size=14, color=LIGHT)

# ---------------------------------------------------------------------------
OUT = 'Presentation.pptx'
prs.save(OUT)
print('saved', OUT, f'({len(prs.slides.__iter__.__self__._sldIdLst)} slides)')
